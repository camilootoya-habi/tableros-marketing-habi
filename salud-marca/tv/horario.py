"""De dónde salen los spots de cada día.

DOS FUENTES, y el reporte usa la mejor disponible para cada fecha:

1. AS-RUN (`spots.csv`) — lo que REALMENTE salió al aire, con hora exacta. Llega semanal de
   la central como xlsx y `ingesta.py` lo convierte a este CSV. Es la fuente de verdad.
2. HORARIO PROYECTADO — para las fechas que todavía no tienen as-run. Se arma con el patrón
   observado: qué canal, qué programa y a qué HORA hubo spot cada día de la semana, según el
   as-run más reciente que sí cubra ese día de semana.

EL DINERO NO VIVE AQUÍ. `spots.csv` es público (este repo lo es) y solo lleva calendario.
La inversión y el TRP van agregados por día y franja en el secret `TV_INVERSION_JSON`, que
`cargar_inversion()` lee. Ver el docstring de `ingesta.py` para el porqué.

POR QUÉ ALCANZA CON PROYECTAR LA HORA Y NO EL MINUTO: el reporte diario usa el estimador de
banda (`estimador.exceso_por_hora`), que solo pregunta qué horas tuvieron spot. Se midió
además cuánta deriva tolera el estimador de minuto: hasta ±15 min sigue siendo significativo
(t=2.02), y a ±20 se cae (t=1.65). O sea que aun proyectando minutos el método aguantaría
desvíos moderados — pero no hace falta correr ese riesgo, porque el estimador de banda no
depende del minuto en absoluto.

LÍMITE CONOCIDO: la proyección asume que la parrilla se repite semana a semana. Hoy hay un
solo as-run (semana 37) así que esa regularidad NO está verificada contra datos — es un
supuesto de negocio, no una medición. Cuando haya 3-4 as-run, `regularidad()` la mide y el
reporte empieza a avisar si el patrón deja de cumplirse.
"""
import csv
import datetime
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))

PESTANA = "Reporte semanal"
COL = {"canal": 4, "fecha": 6, "hora": 7, "programa": 8, "trp": 10, "inversion": 11,
       "franja": 13}


def leer_asrun(ruta):
    """Lee un xlsx de la central → [{ts, canal, programa, trp, inversion, franja}, ...].

    `ts` es datetime al minuto en hora CDMX. El huso quedó verificado: el perfil horario de
    GA4 convertido a America/Mexico_City y el de los registros web tienen los dos su valle a
    las 4am y su pico por la tarde, consistentes con la hora de emisión del as-run.
    """
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    if PESTANA not in wb.sheetnames:
        raise ValueError(f"{os.path.basename(ruta)} no tiene la pestaña '{PESTANA}'")
    ws = wb[PESTANA]
    filas = []
    for fila in ws.iter_rows(min_row=6, values_only=True):
        fecha = fila[COL["fecha"] - 1] if len(fila) >= COL["fecha"] else None
        if not isinstance(fecha, datetime.datetime):
            continue
        crudo = str(fila[COL["hora"] - 1] or "")
        partes = crudo.split(":")
        if len(partes) < 2 or not partes[0].strip().isdigit():
            continue
        hh, mm = int(partes[0]), int(partes[1])
        filas.append({
            "ts": datetime.datetime.combine(fecha.date(), datetime.time(hh, mm)),
            "canal": fila[COL["canal"] - 1],
            "programa": fila[COL["programa"] - 1],
            "trp": float(fila[COL["trp"] - 1] or 0),
            "inversion": float(fila[COL["inversion"] - 1] or 0),
            "franja": fila[COL["franja"] - 1],
            "origen": "as-run",
        })
    return sorted(filas, key=lambda f: f["ts"])


def cargar(ruta=None):
    """Lee `spots.csv` → [{ts, canal, programa, franja, origen}, ...]. Sin dinero ni TRP."""
    ruta = ruta or os.path.join(HERE, "spots.csv")
    if not os.path.exists(ruta):
        return []
    out = []
    with open(ruta, encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            out.append({
                "ts": datetime.datetime.combine(
                    datetime.date.fromisoformat(r["fecha"]),
                    datetime.time(int(r["hora"]), int(r["minuto"]))),
                "canal": r["canal"], "programa": r["programa"],
                "franja": r["franja"] or None, "origen": "as-run"})
    return sorted(out, key=lambda f: f["ts"])


def cargar_inversion():
    """{fecha: {franja: {spots, trp, inversion}}} desde el secret o el archivo local.

    Devuelve {} si no hay ninguno: el reporte entonces publica el calendario y el
    incremental, y omite las cifras de dinero. Es preferible una tarjeta sin inversión a un
    job caído — y a inventar ceros, que se leerían como 'no se invirtió nada'.
    """
    crudo = os.environ.get("TV_INVERSION_JSON", "").strip()
    if crudo:
        try:
            return json.loads(crudo)
        except ValueError:
            print("WARN: TV_INVERSION_JSON no es JSON válido; se ignora")
            return {}
    local = os.path.join(HERE, "inversion.json")
    if os.path.exists(local):
        with open(local, encoding="utf-8") as f:
            return json.load(f)
    return {}


def totales_del_dia(inversion, fecha):
    """(spots, trp, inversion_total, [(franja, inversion, spots), ...]) para una fecha."""
    dia = inversion.get(fecha.isoformat(), {})
    if not dia:
        return None
    por_franja = [(f, v["inversion"], v["spots"]) for f, v in dia.items()]
    return (sum(v["spots"] for v in dia.values()),
            sum(v["trp"] for v in dia.values()),
            sum(v["inversion"] for v in dia.values()),
            sorted(por_franja, key=lambda x: -x[1]))


def _patron(spots):
    """{dia_semana: [{hora, canal, programa, trp, inversion, franja}, ...]} del as-run más
    reciente que cubra ese día de semana. Se toma el más reciente y no un promedio de todos
    porque una parrilla que cambió debe reemplazar a la vieja, no mezclarse con ella."""
    por_dow = {}
    for f in spots:
        dow = f["ts"].weekday()
        fecha = f["ts"].date()
        actual = por_dow.get(dow)
        if actual is None or fecha > actual["fecha"]:
            por_dow[dow] = {"fecha": fecha, "spots": []}
        if fecha == por_dow[dow]["fecha"]:
            por_dow[dow]["spots"].append(f)
    return {dow: v["spots"] for dow, v in por_dow.items()}


def proyectar(spots, fecha):
    """Spots esperados para `fecha` (date) según el patrón del mismo día de semana.

    Devuelve [] si no hay as-run que cubra ese día de semana — el reporte entonces dice que
    no puede estimar, en vez de asumir cero spots, que se leería como 'la TV no salió'.

    NUNCA proyecta hacia ATRÁS del primer as-run. El patrón describe una campaña que empezó
    el 7-sep; aplicarlo a agosto inventaba spots en días sin televisión y contaminaba tanto el
    factor del día como la detección de anomalías en el tramo pre-campaña que el tablero
    dibuja justamente como referencia.
    """
    if spots and fecha < min(f["ts"].date() for f in spots):
        return []
    patron = _patron(spots).get(fecha.weekday(), [])
    out = []
    for f in patron:
        out.append({**f,
                    "ts": datetime.datetime.combine(fecha, f["ts"].time()),
                    "origen": "proyectado"})
    return out


def para_fecha(spots, fecha):
    """As-run si existe para esa fecha; si no, el horario proyectado."""
    reales = [f for f in spots if f["ts"].date() == fecha]
    if reales:
        return reales
    return proyectar(spots, fecha)


def horas_con_spot(spots_dia):
    return {f["ts"].hour for f in spots_dia}


def franja_de_hora(spots_dia):
    """{hora: franja}. Si dos spots de la misma hora traen franja distinta (pasa en los
    bordes de bloque) gana la de mayor jerarquía: AAA es más específica que A."""
    orden = {"AAA": 3, "AA": 2, "A": 1}
    out = {}
    for f in spots_dia:
        h, nueva = f["ts"].hour, f.get("franja")
        if nueva is None:
            continue
        if h not in out or orden.get(nueva, 0) > orden.get(out[h], 0):
            out[h] = nueva
    return out


def regularidad(spots):
    """Qué tan estable es la parrilla entre semanas: fracción de horas-con-spot de cada día
    de semana que se repiten en la semana siguiente.

    Devuelve None con menos de dos semanas de as-run — que es el caso hoy. Existe para que
    el reporte pueda avisar cuando el supuesto de repetición deje de cumplirse, en vez de
    seguir proyectando en silencio sobre una parrilla que cambió.
    """
    por_semana = {}
    for f in spots:
        iso = f["ts"].isocalendar()
        por_semana.setdefault((iso[0], iso[1]), set()).add((f["ts"].weekday(), f["ts"].hour))
    semanas = sorted(por_semana)
    if len(semanas) < 2:
        return None
    coincidencias = []
    for a, b in zip(semanas, semanas[1:]):
        prev, sig = por_semana[a], por_semana[b]
        if not prev:
            continue
        coincidencias.append(len(prev & sig) / len(prev))
    if not coincidencias:
        return None
    return sum(coincidencias) / len(coincidencias)
